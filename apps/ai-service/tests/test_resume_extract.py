"""RI-2 — `app/resume_import/extract.py`.

WHAT THESE TESTS ARE ACTUALLY GUARDING. The extractor is the only thing standing
between a file a stranger uploaded and a decoder, and it is also the thing that decides
what text a language model will later be asked to cite. So the suite is in three parts,
and the middle one is the point:

  1. It extracts what is there — PDF text layers, DOCX bodies, OCR'd scans.
  2. It refuses what it should refuse — a mislabelled file, a password, a page count
     nothing honest has, an image sized to exhaust memory.
  3. It degrades rather than raises, ALWAYS, because ruling D9 says an unreadable
     résumé costs the worker a sentence of Hinglish and never the flow.

EVERY FIXTURE IS BUILT, NOT COMMITTED. The PDFs are assembled here byte by byte, the
DOCX by python-docx, the scans by Pillow. A committed binary fixture is a thing nobody
can review in a diff and nobody can adjust without a hex editor; a generated one states
its own contents in the code that makes it. It also lets each test assert the fixture
CONTAINS the thing before asserting the extractor FOUND it — the vacuity check, written
first, because a fixture that never held the phrase makes the real assertion
unfalsifiable.
"""

from __future__ import annotations

import io
import logging
import os
import sys
import zipfile
from pathlib import Path

import pytest

# `import ... as` rather than `from app.resume_import import extract`: the package
# re-exports the FUNCTION under that name, so the `from` form binds the function and the
# module is unreachable. Both forms are legal and only one of them is the module.
import app.resume_import.extract as extract_mod
from app.resume_import.extract import (
    DEGRADED_REASONS,
    EXTRACTION_METHODS,
    ExtractionLimits,
    extract,
)

PDF_MIME = "application/pdf"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PNG_MIME = "image/png"

# MOST FIXTURES HERE ARE SHORTER THAN A REAL RÉSUMÉ, and `min_text_layer_chars` exists
# precisely to treat a very thin text layer as a scan. Without this override the tidy
# two-line fixtures below would all fall through to the OCR path and these tests would
# quietly be measuring something else. `test_a_text_layer_too_thin_to_be_real_...` keeps
# the production default honest by exercising it directly.
TEXT_LIMITS = ExtractionLimits(min_text_layer_chars=1)

OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


# ===========================================================================
# Fixture builders
# ===========================================================================


def _escape(text: str) -> bytes:
    """PDF literal-string escaping. Our text is tame; the escape is not optional anyway."""
    out = text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")
    return out.encode("latin-1", "replace")


def build_pdf(pages: list[list[str]]) -> bytes:
    """A minimal, VALID PDF whose pages carry a real text layer.

    Hand-assembled with a correct cross-reference table rather than leaning on pypdf's
    repair path: a fixture that only works because the reader is forgiving is not
    evidence about the reader we ship.
    """
    n = len(pages)
    page_ids = [3 + 2 * i for i in range(n)]
    content_ids = [4 + 2 * i for i in range(n)]
    font_id = 3 + 2 * n

    bodies: dict[int, bytes] = {}
    kids = " ".join(f"{pid} 0 R" for pid in page_ids).encode()
    bodies[1] = b"<< /Type /Catalog /Pages 2 0 R >>"
    bodies[2] = b"<< /Type /Pages /Kids [" + kids + b"] /Count " + str(n).encode() + b" >>"
    bodies[font_id] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"

    for index, lines in enumerate(pages):
        stream = b"BT /F1 12 Tf 72 720 Td 14 TL\n"
        for line in lines:
            stream += b"(" + _escape(line) + b") Tj T*\n"
        stream += b"ET"
        bodies[content_ids[index]] = (
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"
        )
        bodies[page_ids[index]] = (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents "
            + str(content_ids[index]).encode()
            + b" 0 R /Resources << /Font << /F1 "
            + str(font_id).encode()
            + b" 0 R >> >> >>"
        )

    out = bytearray(b"%PDF-1.4\n")
    offsets: dict[int, int] = {}
    for obj_id in sorted(bodies):
        offsets[obj_id] = len(out)
        out += str(obj_id).encode() + b" 0 obj\n" + bodies[obj_id] + b"\nendobj\n"

    xref_at = len(out)
    size = max(bodies) + 1
    out += b"xref\n0 " + str(size).encode() + b"\n0000000000 65535 f \n"
    for obj_id in range(1, size):
        out += f"{offsets.get(obj_id, 0):010d} 00000 n \n".encode()
    out += (
        b"trailer\n<< /Size "
        + str(size).encode()
        + b" /Root 1 0 R >>\nstartxref\n"
        + str(xref_at).encode()
        + b"\n%%EOF\n"
    )
    return bytes(out)


def build_docx(
    paragraphs: list[str],
    *,
    table_after: int | None = None,
    table_rows: list[list[str]] | None = None,
    header: str | None = None,
    footer: str | None = None,
) -> bytes:
    """A .docx, optionally with a table INTERLEAVED at a known position."""
    from docx import Document

    document = Document()
    if header:
        document.sections[0].header.paragraphs[0].text = header
    if footer:
        document.sections[0].footer.paragraphs[0].text = footer
    for index, text in enumerate(paragraphs):
        document.add_paragraph(text)
        if table_after is not None and index == table_after and table_rows:
            table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
            for r, row in enumerate(table_rows):
                for c, cell in enumerate(row):
                    table.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def find_font(size: int):
    """A real TrueType face, or None. Tesseract cannot read Pillow's 11px bitmap default."""
    from PIL import ImageFont

    candidates = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            try:
                return ImageFont.truetype(path, size)
            except OSError:
                continue
    return None


def build_png(lines: list[str], *, size: tuple[int, int] = (1400, 500)):
    """A white page with black text — a stand-in for a photographed résumé."""
    from PIL import Image, ImageDraw

    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    font = find_font(64)
    y = 40
    for line in lines:
        draw.text((40, y), line, fill="black", font=font)
        y += 110
    return image


def png_bytes(lines: list[str], **kwargs) -> bytes:
    buffer = io.BytesIO()
    build_png(lines, **kwargs).save(buffer, format="PNG")
    return buffer.getvalue()


def scanned_pdf_bytes(lines: list[str]) -> bytes:
    """A PDF that is ONE RASTER IMAGE and carries no text layer at all — a scan."""
    buffer = io.BytesIO()
    build_png(lines).save(buffer, format="PDF", resolution=150.0)
    return buffer.getvalue()


def encrypted_pdf(user_password: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    source = PdfReader(io.BytesIO(build_pdf([["CNC Turner Pune five years"]])))
    writer = PdfWriter(clone_from=source)
    writer.encrypt(owner_password="owner", user_password=user_password)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# ===========================================================================
# The OCR seam
# ===========================================================================


class FakeTesseract:
    """Stands in for the binary so every decision ABOVE the OCR call is testable here.

    `image_to_data`'s real shape is a dict of parallel lists — one entry per detected
    word, with structural rows carrying `conf == -1`. Reproducing THAT shape, rather
    than a convenient one, is what makes the tests about confidence and line grouping
    mean anything: a mock that returns whatever is asked of it cannot catch a contract
    mismatch.
    """

    class Output:
        DICT = "dict"

    def __init__(self, pages):
        self.pages = list(pages)
        self.calls = 0

    def get_tesseract_version(self):
        return "5.3.0"

    def image_to_data(self, image, lang, output_type):  # noqa: ARG002
        page = self.pages[min(self.calls, len(self.pages) - 1)] if self.pages else []
        self.calls += 1
        data: dict[str, list] = {
            "text": [],
            "conf": [],
            "page_num": [],
            "block_num": [],
            "par_num": [],
            "line_num": [],
        }
        for text, conf, (page_num, block, par, line) in page:
            data["text"].append(text)
            data["conf"].append(conf)
            data["page_num"].append(page_num)
            data["block_num"].append(block)
            data["par_num"].append(par)
            data["line_num"].append(line)
        return data


def one_word_per_line(texts: list[str], conf: float):
    """Each word on its own Tesseract line, all in the same block and paragraph."""
    return [(t, conf, (1, 1, 1, i + 1)) for i, t in enumerate(texts)]


@pytest.fixture
def fake_ocr(monkeypatch):
    """Install a FakeTesseract in place of the binary, and re-probe `ocr_available()`.

    `monkeypatch.setitem(sys.modules, ...)` rather than a plain assignment: a leaked
    fake in `sys.modules` would silently satisfy the real-binary tests further down and
    turn this file's most important coverage into theatre.
    """

    def install(pages):
        fake = FakeTesseract(pages)
        monkeypatch.setitem(sys.modules, "pytesseract", fake)
        extract_mod.ocr_available.cache_clear()
        return fake

    return install


@pytest.fixture(autouse=True)
def _clear_ocr_probe():
    """`ocr_available` is `lru_cache`d because it spawns a process. Tests that swap the
    binary out must not inherit a neighbour's answer."""
    extract_mod.ocr_available.cache_clear()
    yield
    extract_mod.ocr_available.cache_clear()


# ===========================================================================
# 1. The closed vocabularies
# ===========================================================================


def test_the_degraded_vocabulary_is_exactly_this_and_drift_is_a_failure():
    """The Python half of a cross-language contract, pinned against a literal.

    This does NOT prove agreement with `RESUME_IMPORT_FAILURES` in packages/types — that
    pin is owed by RI-3, and `extract.py`'s docstring says why it cannot be written yet
    (the TypeScript constant lands with PR #1478 and is not on `main`, so the pin would
    either fail on `main` or be written to skip, and a skipping pin is worse than none).
    What this does prove is that nobody widens the vocabulary on THIS side without
    touching a test that names the other one.
    """
    assert set(DEGRADED_REASONS) == {
        "no_text_layer",
        "ocr_below_floor",
        "unsupported_document",
        "encrypted_document",
        "empty_document",
        "parse_unavailable",
    }
    assert EXTRACTION_METHODS == ("pdf_text", "docx", "ocr")


def test_a_reason_outside_the_vocabulary_raises_rather_than_reaching_the_database():
    """`worker_resume_import` has a CHECK for every one of these. A typo'd reason that
    got this far would surface as a constraint violation at the INSERT, long after the
    document is gone."""
    with pytest.raises(ValueError, match="closed vocabulary"):
        extract_mod._degraded("ocr_was_a_bit_sad")


# ===========================================================================
# 2. The declared MIME is a claim, not a fact
# ===========================================================================


def test_an_unknown_mime_is_refused_before_any_decoder_sees_the_bytes():
    assert extract(b"%PDF-1.4", mime="text/plain").degraded_reason == "unsupported_document"


def test_an_empty_upload_is_named_empty_not_unsupported():
    assert extract(b"", mime=PDF_MIME).degraded_reason == "empty_document"


def test_a_png_announced_as_a_pdf_never_reaches_pypdf(monkeypatch):
    """THE CLAIM-VS-FACT CHECK. The MIME arrives from the client, and Supabase records
    whatever the upload declared — so the confirm call and the stored object metadata are
    two readings of ONE untrusted source, not two sources. Sniffing the content is the
    only independent one there is.

    ASSERTING ON THE DECODER, NOT ON THE OUTCOME. The reason is a mutation: deleting the
    sniff check entirely left this test GREEN, because pypdf then chokes on the PNG and
    returns the same `unsupported_document` by accident. Identical verdict, no guard —
    and a test that cannot tell those apart is not testing the guard. What the name
    promises is that pypdf never sees the bytes, so that is what is measured.
    """
    import pypdf

    payload = png_bytes(["CNC Turner"])
    assert payload.startswith(b"\x89PNG"), "vacuity: the fixture must actually be a PNG"

    opened: list[str] = []
    original = pypdf.PdfReader
    monkeypatch.setattr(
        pypdf, "PdfReader", lambda *a, **k: (opened.append("read"), original(*a, **k))[1]
    )

    assert extract(payload, mime=PDF_MIME).degraded_reason == "unsupported_document"
    assert opened == [], "the PDF reader was handed a file that is not a PDF"


def test_a_docx_announced_as_an_image_never_reaches_pillow(monkeypatch):
    """The same rule from the other side: a zip must not be opened as an image."""
    from PIL import Image

    payload = build_docx(["CNC Turner"])
    assert payload.startswith(b"PK\x03\x04"), "vacuity: a .docx is a zip"

    opened: list[str] = []
    original = Image.open
    monkeypatch.setattr(
        Image, "open", lambda *a, **k: (opened.append("open"), original(*a, **k))[1]
    )

    assert extract(payload, mime=PNG_MIME).degraded_reason == "unsupported_document"
    assert opened == [], "Pillow was handed a file that is not an image"


def test_a_password_protected_word_file_says_password_not_broken():
    """An encrypted .docx is an OLE container, not a zip. Telling a worker 'this file is
    damaged' when the truth is 'this file has a password' sends them off to re-export a
    document that was fine."""
    assert extract(OLE_MAGIC + b"\x00" * 512, mime=DOCX_MIME).degraded_reason == (
        "encrypted_document"
    )


def test_an_ole_container_announced_as_a_pdf_is_merely_unsupported():
    """The same bytes, a different claim: nothing here says a PDF has a password."""
    assert extract(OLE_MAGIC + b"\x00" * 512, mime=PDF_MIME).degraded_reason == (
        "unsupported_document"
    )


# ===========================================================================
# 3. PDF — the text layer
# ===========================================================================


def test_a_two_page_pdf_yields_contiguous_indices_and_truthful_page_numbers():
    payload = build_pdf([["CNC Turner", "5 years Pune"], ["ITI Fitter 2019"]])
    assert b"CNC Turner" in payload and b"ITI Fitter 2019" in payload, "vacuity"

    result = extract(payload, mime=PDF_MIME, limits=TEXT_LIMITS)

    assert result.degraded_reason is None
    assert result.method == "pdf_text"
    assert result.page_count == 2
    assert result.ocr_confidence is None
    assert [line.index for line in result.lines] == list(range(len(result.lines)))
    assert "CNC Turner" in result.text
    assert "ITI Fitter 2019" in result.text
    # A page number is a locator. If it is wrong it is worse than absent.
    assert all(line.page == 1 for line in result.lines if "CNC Turner" in line.text)
    assert all(line.page == 2 for line in result.lines if "ITI Fitter" in line.text)


def test_a_text_layer_too_thin_to_be_real_is_treated_as_a_scan(monkeypatch):
    """The production default (200 chars), exercised directly. A "printed to PDF" scan
    often carries a page number and a stray header — a text layer in the technical sense
    and nothing a résumé could be parsed from."""
    monkeypatch.setattr(extract_mod, "ocr_available", lambda: False)
    result = extract(build_pdf([["Page 1"]]), mime=PDF_MIME)
    # Fell through past the text layer and into the OCR branch, which is unavailable here.
    assert result.degraded_reason == "parse_unavailable"


def test_an_owner_password_with_no_user_password_still_opens():
    """Cybercafé exporters stamp an owner password on almost everything, and such a file
    is readable by every viewer on earth. Refusing it would reject a large share of the
    documents this feature exists to accept."""
    from pypdf import PdfReader

    payload = encrypted_pdf("")
    assert PdfReader(io.BytesIO(payload)).is_encrypted, "vacuity: this must be encrypted"

    result = extract(payload, mime=PDF_MIME, limits=TEXT_LIMITS)
    assert result.degraded_reason is None
    assert "CNC Turner" in result.text


def test_a_real_user_password_is_reported_as_a_password():
    assert extract(encrypted_pdf("secret"), mime=PDF_MIME).degraded_reason == "encrypted_document"


def test_a_document_with_more_pages_than_any_resume_is_refused_whole():
    """NOT truncated to the first N pages. A 200-page PDF is not a long résumé; feeding
    its opening pages to the model would prefill a worker's profile from someone else's
    document."""
    payload = build_pdf([[f"page {i}"] for i in range(25)])
    result = extract(payload, mime=PDF_MIME, limits=ExtractionLimits(max_pages=20))
    assert result.degraded_reason == "unsupported_document"
    assert result.page_count == 25
    assert result.lines == ()


def test_a_page_count_inside_the_ceiling_is_accepted():
    """Testing what the guard PERMITS. A ceiling nobody can reach passes every "does it
    block a huge file" test and gets deleted the first time it blocks a real one."""
    # Long enough to clear `min_text_layer_chars` on the PRODUCTION default, so this
    # exercises the page ceiling and nothing else.
    payload = build_pdf(
        [
            [f"page {i} of a genuinely long resume", "CNC Turner, VMC Operator, Tata Motors Pune"]
            for i in range(6)
        ]
    )
    result = extract(payload, mime=PDF_MIME, limits=ExtractionLimits(max_pages=20))
    assert result.degraded_reason is None
    assert result.page_count == 6


def test_a_corrupt_pdf_degrades_instead_of_raising():
    result = extract(b"%PDF-1.4\n" + b"\xff" * 400, mime=PDF_MIME)
    assert result.degraded_reason == "unsupported_document"


# ===========================================================================
# 4. DOCX
# ===========================================================================


def test_a_table_keeps_its_place_between_the_paragraphs_that_surround_it():
    """Résumés are routinely laid out as tables. Reading all paragraphs and then all
    cells — which is what the convenient `document.paragraphs` / `document.tables`
    properties give you — shuffles a work history into an order nobody wrote."""
    payload = build_docx(
        ["WORK EXPERIENCE", "EDUCATION"],
        table_after=0,
        table_rows=[["Tata Motors", "2019-2023"]],
    )
    result = extract(payload, mime=DOCX_MIME)

    assert result.degraded_reason is None
    assert result.method == "docx"
    texts = [line.text for line in result.lines]
    assert {"WORK EXPERIENCE", "Tata Motors", "EDUCATION"} <= set(texts)
    assert texts.index("WORK EXPERIENCE") < texts.index("Tata Motors") < texts.index("EDUCATION")


def test_a_docx_reports_no_page_count_and_no_page_numbers():
    """A .docx has no pages until something renders it, and this module renders nothing.
    A guessed page number beside a real quote is a fabricated locator."""
    result = extract(build_docx(["CNC Turner"]), mime=DOCX_MIME)
    assert result.page_count is None
    assert result.lines and all(line.page is None for line in result.lines)


def test_the_header_and_footer_are_read_too():
    """Contact details sit in the header on most templates. A name dropped here is a
    field RI-4 then has to ask the worker for."""
    result = extract(
        build_docx(["EXPERIENCE"], header="Ramesh Kumar +91 98765 43210", footer="Page 1 of 2"),
        mime=DOCX_MIME,
    )
    assert "Ramesh Kumar +91 98765 43210" in result.text
    assert "Page 1 of 2" in result.text


def test_an_empty_docx_is_empty_not_unsupported():
    assert extract(build_docx([]), mime=DOCX_MIME).degraded_reason == "empty_document"


def test_a_zip_that_is_not_a_docx_degrades():
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("hello.txt", "not a word document")
    assert extract(buffer.getvalue(), mime=DOCX_MIME).degraded_reason == "unsupported_document"


# ===========================================================================
# 5. Normalization
# ===========================================================================


def test_normalization_folds_compatibility_forms_and_collapses_runs():
    raw = "CNC\u00a0Turner   \uff11\uff10 years\n\n\n   \nVMC  Operator"
    assert extract_mod._normalize_block(raw) == ["CNC Turner 10 years", "VMC Operator"]


def test_a_bidi_override_is_removed_so_the_stored_text_reads_as_it_displays():
    """U+202E makes rendered text run backwards. Left in, a résumé could show one
    employer to a recruiter while the database, the sheet and every event hold another."""
    lines = extract_mod._normalize_block("Employer: \u202eSROTOM ATAT\u202c")
    assert "\u202e" not in lines[0] and "\u202c" not in lines[0]
    assert lines[0] == "Employer: SROTOM ATAT"


def test_a_form_feed_breaks_the_line_rather_than_gluing_two_together():
    """Deleting a page break would splice the last line of one page onto the first line
    of the next and hand the model a sentence that was never in the document."""
    assert extract_mod._normalize_block("CNC Turner\x0cITI Fitter") == ["CNC Turner", "ITI Fitter"]


def test_a_zero_width_joiner_is_removed():
    assert extract_mod._normalize_block("Tata\u200bMotors") == ["TataMotors"]


def test_blank_lines_are_dropped_before_numbering_so_every_index_addresses_something():
    payload = build_pdf([["CNC Turner", "", "   ", "VMC Operator"]])
    result = extract(payload, mime=PDF_MIME, limits=TEXT_LIMITS)
    assert all(line.text.strip() for line in result.lines)
    assert [line.index for line in result.lines] == list(range(len(result.lines)))


# ===========================================================================
# 6. The caps, and saying so out loud
# ===========================================================================


def test_hitting_the_line_cap_is_reported_rather_than_looking_complete():
    payload = build_pdf([[f"line {i}" for i in range(50)]])
    result = extract(
        payload, mime=PDF_MIME, limits=ExtractionLimits(max_lines=10, min_text_layer_chars=1)
    )
    assert len(result.lines) == 10
    assert result.truncated is True


def test_hitting_the_character_cap_is_reported():
    payload = build_pdf([["abcdefghij" for _ in range(50)]])
    result = extract(
        payload, mime=PDF_MIME, limits=ExtractionLimits(max_total_chars=25, min_text_layer_chars=1)
    )
    assert result.truncated is True
    assert sum(len(line.text) for line in result.lines) <= 25


def test_a_pathological_line_is_wrapped_not_dropped():
    """Citation granularity is the point of a line, and a 40,000-character 'line' is not
    a citation — but its content is still the worker's, so it is split, never discarded."""
    payload = build_pdf([["X" * 250]])
    result = extract(
        payload, mime=PDF_MIME, limits=ExtractionLimits(max_line_chars=100, min_text_layer_chars=1)
    )
    assert all(len(line.text) <= 100 for line in result.lines)
    assert "".join(line.text for line in result.lines).count("X") == 250


def test_a_document_that_fits_is_not_marked_truncated():
    """The other half of the rule: `truncated` must not become decorative."""
    result = extract(build_pdf([["CNC Turner Pune"]]), mime=PDF_MIME, limits=TEXT_LIMITS)
    assert result.truncated is False


# ===========================================================================
# 7. OCR — every decision above the binary
# ===========================================================================


def test_a_scan_with_ocr_switched_off_is_named_no_text_layer(monkeypatch):
    """The ops kill-switch, and the only route to `no_text_layer`. Turning OCR off must
    produce a named reason, not a confusing silence."""
    monkeypatch.setattr(extract_mod, "ocr_available", lambda: True)
    result = extract(
        scanned_pdf_bytes(["CNC TURNER"]),
        mime=PDF_MIME,
        limits=ExtractionLimits(ocr_enabled=False),
    )
    assert result.degraded_reason == "no_text_layer"


def test_a_missing_tesseract_is_an_OPS_failure_and_is_named_as_one(monkeypatch, caplog):
    """`no_text_layer` would be true of the document and false about us. Every scanned
    résumé would then fail with a reason pointing the reader at the worker's file instead
    of at a container missing its binary — a feature that looks green and never works."""
    monkeypatch.setattr(extract_mod, "ocr_available", lambda: False)
    with caplog.at_level(logging.WARNING):
        result = extract(scanned_pdf_bytes(["CNC TURNER"]), mime=PDF_MIME)
    assert result.degraded_reason == "parse_unavailable"
    assert any("ocr_unavailable" in record.getMessage() for record in caplog.records)


def test_a_photograph_with_no_ocr_available_degrades_without_raising(monkeypatch):
    monkeypatch.setattr(extract_mod, "ocr_available", lambda: False)
    result = extract(png_bytes(["CNC TURNER"]), mime=PNG_MIME)
    assert result.degraded_reason == "parse_unavailable"
    assert result.page_count == 1


def test_ocr_above_the_floor_produces_lines_a_confidence_and_a_method(fake_ocr):
    fake_ocr([one_word_per_line(["CNC", "TURNER"], 92.0)])
    result = extract(png_bytes(["CNC TURNER"]), mime=PNG_MIME)
    assert result.degraded_reason is None
    assert result.method == "ocr"
    assert result.page_count == 1
    assert result.ocr_confidence == pytest.approx(0.92)
    assert result.text == "CNC\nTURNER"


def test_ocr_below_the_floor_KEEPS_the_score_it_failed_on(fake_ocr):
    """RI-7 has to answer "how often, and by how much, does OCR fall short?" before
    anyone claims this feature saves a worker time. That is unanswerable if the failures
    throw away the only measurement they produced."""
    fake_ocr([one_word_per_line(["CNC", "TURNER"], 31.0)])
    result = extract(png_bytes(["CNC TURNER"]), mime=PNG_MIME)
    assert result.degraded_reason == "ocr_below_floor"
    assert result.method == "ocr"
    assert result.ocr_confidence == pytest.approx(0.31)
    assert result.lines == ()


def test_ocr_that_recovers_nothing_is_an_empty_document(fake_ocr):
    fake_ocr([[]])
    result = extract(png_bytes(["CNC TURNER"]), mime=PNG_MIME)
    assert result.degraded_reason == "empty_document"
    assert result.method == "ocr"


def test_a_scanned_pdf_is_rasterized_and_reaches_the_ocr_seam(fake_ocr):
    """The pypdfium2 leg with the binary stubbed: a PDF carrying no text layer must be
    rendered to an image and offered to OCR, not reported as unreadable."""
    fake = fake_ocr([one_word_per_line(["CNC", "TURNER", "TATA", "MOTORS"], 88.0)])
    result = extract(
        scanned_pdf_bytes(["CNC TURNER"]),
        mime=PDF_MIME,
        limits=ExtractionLimits(ocr_render_dpi=72),
    )
    assert fake.calls == 1, "the rasterizer must have produced exactly one page image"
    assert result.method == "ocr"
    assert result.page_count == 1
    assert "TATA" in result.text


def test_the_confidence_is_weighted_by_LINE_COUNT_across_pages(monkeypatch):
    """A title page of four words must not outvote a dense page of work history."""
    calls = iter([(["A"], 0.40), (["B"] * 9, 0.90)])
    monkeypatch.setattr(extract_mod, "_run_tesseract", lambda image, languages: next(calls))
    per_page, confidence = extract_mod._ocr_pages([object(), object()], ExtractionLimits())
    assert confidence == pytest.approx((0.40 * 1 + 0.90 * 9) / 10)
    assert per_page == [["A"], ["B"] * 9]


def test_structural_rows_and_string_confidences_never_drag_the_score_down(monkeypatch):
    """Tesseract emits `conf == -1` for rows that carry no word, and some builds hand the
    confidence back as a string. Averaging either in makes the floor meaningless."""
    fake = FakeTesseract(
        [
            [
                ("", -1.0, (1, 1, 0, 0)),
                ("CNC", "90", (1, 1, 1, 1)),
                ("block", -1.0, (1, 2, 0, 0)),
                ("TURNER", 90.0, (1, 1, 1, 1)),
            ]
        ]
    )
    monkeypatch.setitem(sys.modules, "pytesseract", fake)
    lines, confidence = extract_mod._run_tesseract(object(), languages="eng")
    assert confidence == pytest.approx(0.90)
    assert lines == ["CNC TURNER"]


def test_words_are_grouped_into_lines_in_tesseracts_own_reading_order(monkeypatch):
    fake = FakeTesseract(
        [
            [
                ("Operator", 90.0, (1, 1, 1, 2)),
                ("CNC", 90.0, (1, 1, 1, 1)),
                ("Turner", 90.0, (1, 1, 1, 1)),
                ("VMC", 90.0, (1, 1, 1, 2)),
            ]
        ]
    )
    monkeypatch.setitem(sys.modules, "pytesseract", fake)
    lines, _ = extract_mod._run_tesseract(object(), languages="eng")
    assert lines == ["CNC Turner", "Operator VMC"]


# ===========================================================================
# 8. The decompression bomb
# ===========================================================================


def test_an_oversized_image_is_refused_from_its_HEADER_before_any_decode(monkeypatch, fake_ocr):
    """A ~1 KB PNG can declare 60000x60000 and ask Pillow for ~14 GB. `Image.open` is
    lazy, so the guard has to run between the header and the decode — checking after
    `convert()` would be checking after the allocation that was the whole danger."""
    fake_ocr([one_word_per_line(["CNC"], 95.0)])
    # Built BEFORE the spy is installed: Pillow's own drawing and encoding must not be
    # mistaken for the decode this test is watching for.
    payload = png_bytes(["CNC"], size=(400, 400))

    from PIL import Image

    decoded: list[str] = []
    original = Image.Image.convert

    def spy(self, *args, **kwargs):
        decoded.append("decoded")
        return original(self, *args, **kwargs)

    monkeypatch.setattr(Image.Image, "convert", spy)

    result = extract(payload, mime=PNG_MIME, limits=ExtractionLimits(max_image_pixels=1000))

    assert result.degraded_reason == "unsupported_document"
    assert decoded == [], "the image was decoded despite exceeding the pixel ceiling"


def test_an_image_inside_the_ceiling_is_still_processed(fake_ocr):
    """Testing what the guard PERMITS. An over-broad ceiling passes every "does it block
    a bomb" test and gets deleted the first time it blocks a real photograph."""
    fake_ocr([one_word_per_line(["CNC"], 95.0)])
    payload = png_bytes(["CNC"], size=(400, 400))
    result = extract(payload, mime=PNG_MIME, limits=ExtractionLimits(max_image_pixels=1_000_000))
    assert result.degraded_reason is None
    assert result.method == "ocr"


# ===========================================================================
# 9. Privacy
# ===========================================================================


def test_not_one_character_of_the_document_reaches_the_log(caplog):
    """CLAUDE.md §2 for logs, which ADR-0041 §3 did NOT move. D5 widened what the MODEL
    may see and said nothing about what a log may hold; they were never one question."""
    secret = "Ramesh Kumar Tata Motors 9876543210"
    payload = build_pdf([[secret, "PAN ABCDE1234F"]])
    with caplog.at_level(logging.DEBUG):
        result = extract(payload, mime=PDF_MIME, limits=TEXT_LIMITS)

    assert secret in result.text, "vacuity: the fixture text must actually have been extracted"
    emitted = "\n".join(
        record.getMessage() + repr(getattr(record, "extra", "")) for record in caplog.records
    )
    for fragment in ("Ramesh", "Tata Motors", "9876543210", "ABCDE1234F"):
        assert fragment not in emitted


def test_the_one_log_line_carries_counts_and_nothing_else(caplog):
    with caplog.at_level(logging.INFO):
        extract(build_pdf([["CNC Turner Pune"]]), mime=PDF_MIME, limits=TEXT_LIMITS)

    records = [r for r in caplog.records if r.getMessage() == "resume_import.extracted"]
    assert len(records) == 1
    payload = getattr(records[0], "extra", None)
    # The formatter reads `record.extra`; a FLAT `extra={...}` is accepted by the stdlib
    # and then dropped on the floor by JsonFormatter — structured in the source, empty in
    # the log. This assertion is what would catch that.
    assert payload is not None
    assert set(payload) == {
        "method",
        "line_count",
        "page_count",
        "ocr_confidence",
        "degraded_reason",
        "truncated",
    }


# ===========================================================================
# 10. The real binary — and a guard against this section silently never running
# ===========================================================================


def test_the_ocr_leg_is_actually_exercised_in_ci():
    """THE ANTI-VACUITY GUARD FOR THIS WHOLE SECTION.

    Every test below is skipped unless a real Tesseract is present. On a developer laptop
    that is correct; in CI it would be a lie — the suite would go green having never once
    run the binary, and a broken Dockerfile or a dropped apt package would ship unnoticed.
    A skip is only legal where nobody was promised the coverage.
    """
    if not os.environ.get("CI"):
        pytest.skip("local box: a missing Tesseract is expected here and is not a failure")
    assert extract_mod.ocr_available(), (
        "CI must install tesseract-ocr (see the ai-service job in .github/workflows/ci.yml). "
        "Without it every OCR test below skips and the suite reports green on nothing."
    )
    assert find_font(48) is not None, (
        "CI must install a TrueType font (fonts-dejavu-core). Pillow's 11px bitmap default "
        "is not legible to Tesseract, so the OCR fixtures would be unreadable by design."
    )


needs_ocr = pytest.mark.skipif(
    not extract_mod.ocr_available() or find_font(48) is None,
    reason="needs a real tesseract binary and a TrueType font",
)


@needs_ocr
def test_real_tesseract_reads_a_photographed_resume():
    result = extract(png_bytes(["CNC TURNER", "TATA MOTORS"]), mime=PNG_MIME)
    assert result.degraded_reason is None, result.degraded_reason
    assert result.method == "ocr"
    assert result.ocr_confidence is not None and result.ocr_confidence > 0.6
    assert "CNC" in result.text.upper()


@needs_ocr
def test_real_tesseract_reads_a_scanned_pdf_through_the_rasterizer():
    """The pypdfium2 leg, end to end: a PDF whose only content is a raster image."""
    result = extract(scanned_pdf_bytes(["CNC TURNER"]), mime=PDF_MIME)
    assert result.degraded_reason is None, result.degraded_reason
    assert result.method == "ocr"
    assert result.page_count == 1
    assert "CNC" in result.text.upper()
