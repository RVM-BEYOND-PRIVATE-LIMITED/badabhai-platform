"""Uploaded document → numbered lines of text. Deterministic, offline, no LLM (RI-2).

WHAT THIS IS FOR. ADR-0041 lets a worker hand us a résumé instead of answering the
whole interview from scratch. Everything downstream of this module is a language
model reading text; this module is the only thing that decides WHAT TEXT THERE IS.
That split is load-bearing rather than tidy:

  * RI-3's prompt contract is "cite a `line_index` and quote it character for
    character, or return null". That is only checkable because the lines the model
    cites are the lines produced HERE, by a pure function, stored alongside the
    parse. A model that invents a line index is caught by comparison, not trust.

  * The AI boundary stays TEXT-ONLY. `Message = dict[str, str]` in `app/ai/router.py`
    has no multimodal content block and gains none: a photographed résumé is OCR'd
    locally (Tesseract, in this container) and only the recovered characters travel.
    Ruling D3 accepts photographs precisely BECAUSE no new sub-processor is involved
    — swapping this for a cloud vision API would silently reopen a decision the owner
    already took.

DEGRADES, NEVER RAISES. Every path returns an `ExtractionResult`. A corrupt file, a
password, a 400-page scan and a missing Tesseract binary all come back as a
`degraded_reason` from a CLOSED vocabulary, never as an exception and never as a
half-filled result. Ruling D9: an unreadable file must say so plainly and drop the
worker into the ordinary Hinglish flow — it is never a dead end, so there is nothing
for a caller to catch.

PRIVACY. Document bytes, extracted text and object keys are NEVER logged. The one log
line this module emits carries a method name, counts and a confidence — the same
discipline as `app/storage.py` and `app/audio_chunk.py`. Note that this is stricter
than ADR-0041 §3 requires of the MODEL (which now sees the résumé fully unmasked):
what a prompt may carry and what a log may carry were never the same question, and
D5 moved only the first one.

UNTRUSTED INPUT, AND THE BOUNDS THAT FOLLOW. These bytes came off a worker's phone
through a signed upload URL. Every parser below is a decoder pointed at a file we did
not write, so each limit in `ExtractionLimits` is a guard rather than a preference:

  * `max_pages` — a 10 MiB PDF can declare thousands of pages, and OCR is seconds per
    page. Without this, one upload occupies the service for an hour.
  * `max_image_pixels` — the decompression bomb. A ~1 KB PNG can declare 60000x60000
    and ask Pillow for ~14 GB. Checked from the HEADER, before any decode.
  * `max_total_chars` / `max_lines` — these lines become RI-3's prompt. An unbounded
    document is an unbounded prompt, which is unbounded spend against the TD27 caps.
  * `max_line_chars` — a PDF content stream can emit one "line" the length of the
    file. Long lines are hard-wrapped rather than dropped: citation granularity is
    the point, and a 40,000-character line is not a citation.

A result that hit one of these caps says so (`truncated`) instead of quietly looking
complete. A silent cap reads as "we processed the whole document" when we did not.
"""

from __future__ import annotations

import functools
import io
import re
import unicodedata
from dataclasses import dataclass
from typing import Final

from ..logging_config import get_logger

logger = get_logger("ai-service.resume_import")

# ---------------------------------------------------------------------------
# The closed vocabularies. Both MIRROR TypeScript, and the mirror is deliberate.
# ---------------------------------------------------------------------------

# `RESUME_EXTRACTION_METHODS` in packages/types/src/index.ts, character for character.
EXTRACTION_METHODS: Final = ("pdf_text", "docx", "ocr")

# The subset of `RESUME_IMPORT_FAILURES` (packages/types) that EXTRACTION can produce.
# The three values missing from this tuple — `parse_deadline_exceeded`,
# `parse_output_invalid` and the LLM half of `parse_unavailable` — belong to RI-3 and
# cannot arise here, because nothing here calls a model.
#
# THE CROSS-LANGUAGE PIN IS OWED, AND IT IS OWED TO RI-3. `tests/test_stt.py` shows the
# shape it must take: read the TypeScript, regex the constant, compare — never restate
# the strings in Python, because a test that hard-codes both sides passes while somebody
# changes one of them. It cannot be written yet: `RESUME_IMPORT_FAILURES` lands with
# RI-1 (PR #1478) and is not on `main` at the time this file is written, so the pin would
# either fail on `main` or be written to skip, and a skipping pin is worse than none.
# `test_resume_extract.py` therefore pins the Python side against a literal set, which
# catches drift on this side; RI-3 closes the loop across the boundary.
DEGRADED_REASONS: Final = (
    "no_text_layer",
    "ocr_below_floor",
    "unsupported_document",
    "encrypted_document",
    "empty_document",
    "parse_unavailable",
)

# ---------------------------------------------------------------------------
# Limits
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class ExtractionLimits:
    """Bounds on an untrusted document. See the module docstring for why each exists."""

    # A résumé is one to four pages. Twenty is far past anything honest and still
    # bounds the worst-case OCR run (~20 x a few seconds) inside a parse deadline.
    max_pages: int = 20

    # Pillow's own bomb guard sits near 178 Mpx; this is checked earlier and lower,
    # from the image header, so nothing large is ever decoded. A 300 DPI A4 scan is
    # ~8.7 Mpx, so 40 Mpx leaves room for an oversized phone photo and no more.
    max_image_pixels: int = 40_000_000

    # These lines become a prompt. ~120k characters is far more résumé than exists
    # and still a bounded spend.
    max_total_chars: int = 120_000
    max_lines: int = 4_000
    max_line_chars: int = 1_000

    # Below this many characters a PDF's text layer is treated as absent — page
    # numbers and a stray header, i.e. a scan that was "printed to PDF". PROVISIONAL:
    # RI-7 measures the real distribution over a résumé corpus and calibrates it.
    # Nothing about this number is measured yet, and it must not be treated as if it
    # were (see docs/decisions/0041 §5 and the occupation-floor precedent).
    min_text_layer_chars: int = 200

    # Mean Tesseract word confidence, 0..1, below which the recovered text is treated
    # as unusable. ALSO PROVISIONAL and also RI-7's to calibrate. Ruling D9 is what
    # makes a floor safe to guess at all: falling below it costs the worker a
    # sentence of Hinglish, not the flow.
    ocr_confidence_floor: float = 0.60

    # Rasterization resolution for a scanned PDF. Tesseract's own guidance is ~300 DPI
    # for body text; PDF user space is 72 DPI, hence the scale factor downstream.
    ocr_render_dpi: int = 300

    # `eng+hin` — a résumé written in English by a Hindi speaker routinely carries
    # both scripts. Both traineddata files are installed by the Dockerfile; if either
    # is missing Tesseract errors and this degrades to `parse_unavailable`, loudly.
    ocr_languages: str = "eng+hin"

    # THE OPS KILL-SWITCH, and the only way `no_text_layer` is reachable. OCR is the
    # expensive leg; turning it off must degrade a scan to a named reason rather than
    # to a confusing silence.
    ocr_enabled: bool = True


DEFAULT_LIMITS: Final = ExtractionLimits()

# ---------------------------------------------------------------------------
# Results
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Line:
    """One citable line.

    `index` is the position in `ExtractionResult.lines` and is what RI-3's prompt asks
    the model to cite. It is contiguous over the lines actually returned — blank lines
    are dropped BEFORE numbering, so an index always addresses something.

    `page` is 1-based, and is `None` for DOCX, which has no pages until something
    renders it. Guessing a page number there would put a fabricated locator next to a
    real quote, which is precisely the confusion the citation contract exists to stop.
    """

    index: int
    page: int | None
    text: str


@dataclass(frozen=True)
class ExtractionResult:
    """What came out, and — if nothing useful did — the closed reason why.

    `method`, `page_count` and `ocr_confidence` are shaped to land directly in
    `worker_resume_import` (packages/db/src/schema/resume-import.ts), whose CHECK
    constraints already encode the invariants asserted here: a confidence exists only
    for `ocr`, and a page count is either absent or positive.
    """

    method: str | None
    lines: tuple[Line, ...]
    page_count: int | None
    ocr_confidence: float | None
    degraded_reason: str | None
    truncated: bool = False

    @property
    def ok(self) -> bool:
        return self.degraded_reason is None and bool(self.lines)

    @property
    def text(self) -> str:
        """The lines rejoined — what RI-3 hands the model, alongside their indices."""
        return "\n".join(line.text for line in self.lines)


def _degraded(
    reason: str,
    *,
    method: str | None = None,
    page_count: int | None = None,
    ocr_confidence: float | None = None,
) -> ExtractionResult:
    """The ONLY constructor for a failed extraction, so the vocabulary cannot be widened
    by a typo at a call site. A reason outside `DEGRADED_REASONS` is a programming error
    here, not an input error, and raising is correct: the alternative is a value that no
    CHECK constraint in `worker_resume_import` will accept, discovered at the INSERT.

    Not an `assert` — `python -O` strips those, and a guard that disappears under a flag
    the deployment might one day set is not a guard.
    """
    if reason not in DEGRADED_REASONS:
        raise ValueError(f"reason outside the closed vocabulary: {reason!r}")
    return ExtractionResult(
        method=method,
        lines=(),
        page_count=page_count,
        ocr_confidence=ocr_confidence,
        degraded_reason=reason,
    )


# ---------------------------------------------------------------------------
# Content sniffing — the declared MIME is a claim, not a fact
# ---------------------------------------------------------------------------

_PDF_MAGIC = b"%PDF-"
_ZIP_MAGIC = b"PK\x03\x04"
_PNG_MAGIC = b"\x89PNG\r\n\x1a\n"
_JPEG_MAGIC = b"\xff\xd8\xff"
# The OLE2 compound-file header. An ENCRYPTED .docx is not a zip at all — Office wraps
# the real package in an OLE container — so this is how a password-protected Word file
# is told apart from a corrupt one, and it is the difference between telling a worker
# "this file needs a password" and telling them "this file is broken".
_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"

# Where a `%PDF-` header may appear. The spec tolerates leading bytes before the header
# and real-world writers emit them; a bounded scan accepts those without accepting a
# file that merely mentions "%PDF-" somewhere in its body.
_PDF_HEADER_WINDOW = 1024

_MIME_TO_KIND: Final = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "image/jpeg": "image",
    "image/png": "image",
}


def _sniff(data: bytes) -> str | None:
    """The kind this file ACTUALLY is: `pdf`, `docx`, `image`, `ole`, or None."""
    if data.startswith(_OLE_MAGIC):
        return "ole"
    if data.startswith(_PNG_MAGIC) or data.startswith(_JPEG_MAGIC):
        return "image"
    if data.startswith(_ZIP_MAGIC):
        return "docx"
    if _PDF_MAGIC in data[:_PDF_HEADER_WINDOW]:
        return "pdf"
    return None


# ---------------------------------------------------------------------------
# Text normalization
# ---------------------------------------------------------------------------

#
# EXPRESSED AS CODE POINTS, never as the characters themselves. A literal zero-width
# joiner inside a pattern is invisible in every editor, diff and review that would ever
# look at it — the only form of this rule a human can actually check is a list of
# numbers. (The first draft of this file wrote them as escapes and a tool in the chain
# ate a backslash level, silently substituting the very characters being removed.)

# Folded to a single U+0020: tab, no-break space, ogham space mark, the en/em quad
# family, and the narrow, medium and ideographic spaces.
_SPACE_CODEPOINTS: Final = (0x09, 0xA0, 0x1680, *range(0x2000, 0x200B), 0x202F, 0x205F, 0x3000)

# Turned INTO a line break rather than deleted. A vertical tab, a form feed or a Unicode
# line/paragraph separator is where the document said one line ended; deleting it would
# glue two unrelated lines together and produce a citation pointing at a sentence that
# never existed.
_BREAK_CODEPOINTS: Final = (0x0B, 0x0C, 0x85, 0x2028, 0x2029)

# Deleted outright: the C0/C1 controls, the zero-width marks, the BOM — and the
# bidirectional OVERRIDES, which are a security control rather than tidying. U+202E and
# its family make rendered text read differently from the stored text, so a résumé could
# show one employer to a recruiter while the database, the sheet and every event hold
# another. Removing them makes what we store and what a human sees the same string.
_DELETE_CODEPOINTS: Final = tuple(
    codepoint
    for codepoint in (
        *range(0x00, 0x20),
        *range(0x7F, 0xA0),
        *range(0x200B, 0x2010),
        *range(0x202A, 0x202F),
        *range(0x2066, 0x206A),
        0xFEFF,
    )
    # \n survives; \r is already gone by the time this runs.
    if codepoint != 0x0A and codepoint not in _SPACE_CODEPOINTS + _BREAK_CODEPOINTS
)

# ONE table, ONE pass. Built in this order so a code point can never be listed as both
# deleted and replaced without the later entry winning silently — the comprehension
# above removes the overlap at the source instead.
_TRANSLATE_TABLE: Final = {
    **dict.fromkeys(_DELETE_CODEPOINTS),
    **dict.fromkeys(_SPACE_CODEPOINTS, " "),
    **dict.fromkeys(_BREAK_CODEPOINTS, "\n"),
}

_RUNS_RE = re.compile(r" {2,}")


def _normalize_block(raw: str) -> list[str]:
    """One extracted blob → clean, non-empty candidate lines (not yet numbered)."""
    text = raw.replace("\r\n", "\n").replace("\r", "\n")
    # NFKC folds the ligatures and full-width forms PDF writers emit (ﬁ → fi, １ → 1)
    # so a quote of ours matches a quote of the worker's. It is applied BEFORE the
    # lines are stored, so the model still quotes exactly what it was given.
    text = unicodedata.normalize("NFKC", text)
    text = text.translate(_TRANSLATE_TABLE)
    out: list[str] = []
    for piece in text.split("\n"):
        cleaned = _RUNS_RE.sub(" ", piece).strip()
        if cleaned:
            out.append(cleaned)
    return out


def _number(
    blocks: list[tuple[int | None, list[str]]], limits: ExtractionLimits
) -> tuple[tuple[Line, ...], bool]:
    """Assign contiguous indices, applying the line/char caps. Returns (lines, truncated)."""
    lines: list[Line] = []
    total = 0
    truncated = False
    for page, texts in blocks:
        for text in texts:
            # Hard-wrap rather than drop: a pathological line still carries content, and
            # a citation is only useful if it points at something a human can find.
            for start in range(0, len(text), limits.max_line_chars):
                chunk = text[start : start + limits.max_line_chars]
                if len(lines) >= limits.max_lines or total + len(chunk) > limits.max_total_chars:
                    return tuple(lines), True
                lines.append(Line(index=len(lines), page=page, text=chunk))
                total += len(chunk)
    return tuple(lines), truncated


# ---------------------------------------------------------------------------
# OCR
# ---------------------------------------------------------------------------


@functools.lru_cache(maxsize=1)
def ocr_available() -> bool:
    """Is a usable Tesseract binary on this box?

    Cached because it spawns a process. `ocr_available.cache_clear()` is the seam tests
    use, and the only supported way to re-probe.
    """
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
    except Exception:  # noqa: BLE001 — any failure here means "no OCR", never a crash
        return False
    return True


def _run_tesseract(image: object, *, languages: str) -> tuple[list[str], float]:
    """One image → (lines, mean word confidence 0..1).

    THE SEAM. This is the only function in the package that touches the Tesseract
    binary, so it is the only thing a test has to replace to exercise every decision
    above it on a box with no OCR installed. `image_to_data` rather than
    `image_to_string` because it yields the confidence and the line grouping in ONE
    pass — asking twice would let the text and the score describe different runs.
    """
    import pytesseract

    data = pytesseract.image_to_data(
        image, lang=languages, output_type=pytesseract.Output.DICT
    )
    groups: dict[tuple[int, int, int, int], list[str]] = {}
    confidences: list[float] = []
    for i, word in enumerate(data["text"]):
        text = (word or "").strip()
        if not text:
            continue
        try:
            conf = float(data["conf"][i])
        except (TypeError, ValueError):
            conf = -1.0
        # Tesseract reports -1 for structural rows that carry no word. Averaging those
        # in would drag every score toward zero and make the floor meaningless.
        if conf < 0:
            continue
        confidences.append(conf)
        key = (
            int(data["page_num"][i]),
            int(data["block_num"][i]),
            int(data["par_num"][i]),
            int(data["line_num"][i]),
        )
        groups.setdefault(key, []).append(text)

    # Sorting by the composite key IS reading order as Tesseract resolved it — page,
    # then block, then paragraph, then line.
    lines = [" ".join(words) for _, words in sorted(groups.items())]
    mean = (sum(confidences) / len(confidences) / 100.0) if confidences else 0.0
    return lines, mean


def _ocr_pages(images: list[object], limits: ExtractionLimits) -> tuple[list[list[str]], float]:
    """OCR each page image. Returns (per-page lines, mean confidence over all pages)."""
    per_page: list[list[str]] = []
    weighted = 0.0
    counted = 0
    for image in images:
        lines, confidence = _run_tesseract(image, languages=limits.ocr_languages)
        per_page.append(lines)
        if lines:
            # Weighted by line count so a title page of four words does not carry the
            # same weight as a dense page of experience.
            weighted += confidence * len(lines)
            counted += len(lines)
    return per_page, (weighted / counted if counted else 0.0)


# ---------------------------------------------------------------------------
# Per-format extraction
# ---------------------------------------------------------------------------


def _extract_pdf(data: bytes, limits: ExtractionLimits) -> ExtractionResult:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
    except Exception:  # noqa: BLE001 — pypdf raises a wide family on malformed input
        return _degraded("unsupported_document")

    if reader.is_encrypted:
        # An OWNER password with an empty user password is common on résumés exported by
        # cybercafé software, and such a file is perfectly readable. Try that before
        # telling a worker their document needs a password it may not have.
        try:
            opened = reader.decrypt("")
        except Exception:  # noqa: BLE001
            opened = 0
        if not opened:
            return _degraded("encrypted_document")

    try:
        page_count = len(reader.pages)
    except Exception:  # noqa: BLE001
        return _degraded("unsupported_document")

    if page_count == 0:
        return _degraded("empty_document")
    if page_count > limits.max_pages:
        # Deliberately NOT truncated to the first N pages. A 200-page document is not a
        # résumé that happens to be long; treating it as one would hand the model the
        # opening pages of something else entirely and prefill a profile from it.
        return _degraded("unsupported_document", page_count=page_count)

    blocks: list[tuple[int | None, list[str]]] = []
    layer_chars = 0
    for number, page in enumerate(reader.pages, start=1):
        try:
            raw = page.extract_text() or ""
        except Exception:  # noqa: BLE001 — one bad page must not lose the other nineteen
            raw = ""
        texts = _normalize_block(raw)
        layer_chars += sum(len(t) for t in texts)
        blocks.append((number, texts))

    if layer_chars >= limits.min_text_layer_chars:
        lines, truncated = _number(blocks, limits)
        if not lines:
            return _degraded("empty_document", page_count=page_count)
        return ExtractionResult(
            method="pdf_text",
            lines=lines,
            page_count=page_count,
            ocr_confidence=None,
            degraded_reason=None,
            truncated=truncated,
        )

    # No usable text layer: this is a scan or a photograph wrapped in a PDF.
    if not limits.ocr_enabled:
        return _degraded("no_text_layer", page_count=page_count)
    if not ocr_available():
        # AN OPS FAILURE, NAMED AS ONE. Reporting `no_text_layer` here would be true of
        # the document and false about us, and every scanned résumé would fail with a
        # reason that points the reader at the worker's file instead of at a container
        # missing its Tesseract binary.
        # `extra={"extra": {...}}` — the DOUBLE key is the contract JsonFormatter
        # actually implements (app/logging_config.py reads `record.extra`). A flat
        # `extra={...}` is accepted by the stdlib, attaches the keys to the record, and
        # is then dropped on the floor by the formatter: a log line that looks
        # structured in the source and arrives empty.
        logger.warning(
            "resume_import.ocr_unavailable",
            extra={"extra": {"stage": "pdf", "page_count": page_count}},
        )
        return _degraded("parse_unavailable", page_count=page_count)

    try:
        images = _render_pdf_pages(data, page_count, limits)
    except Exception:  # noqa: BLE001
        return _degraded("unsupported_document", page_count=page_count)

    per_page, confidence = _ocr_pages(images, limits)
    return _finish_ocr(
        [(i + 1, texts) for i, texts in enumerate(_normalize_pages(per_page))],
        page_count=page_count,
        confidence=confidence,
        limits=limits,
    )


def _render_pdf_pages(data: bytes, page_count: int, limits: ExtractionLimits) -> list[object]:
    """Rasterize every page. pypdfium2 (Apache-2.0/BSD-3) — NOT PyMuPDF.

    The plan named PyMuPDF; it is AGPL-3.0 unless a commercial licence is bought from
    Artifex, which is not a dependency a closed-source product can take on quietly.
    pypdfium2 wraps the same PDFium engine Chrome renders with, ships manylinux wheels
    (so the Dockerfile's "no compiler in the runtime" rule holds), and needs no system
    package at all.
    """
    import pypdfium2 as pdfium

    scale = limits.ocr_render_dpi / 72.0
    document = pdfium.PdfDocument(data)
    images: list[object] = []
    try:
        for index in range(min(page_count, limits.max_pages)):
            page = document[index]
            bitmap = page.render(scale=scale)
            images.append(bitmap.to_pil())
    finally:
        document.close()
    return images


def _extract_image(data: bytes, limits: ExtractionLimits) -> ExtractionResult:
    if not limits.ocr_enabled:
        return _degraded("no_text_layer", page_count=1)
    if not ocr_available():
        logger.warning("resume_import.ocr_unavailable", extra={"extra": {"stage": "image"}})
        return _degraded("parse_unavailable", page_count=1)

    from PIL import Image

    try:
        image = Image.open(io.BytesIO(data))
    except Exception:  # noqa: BLE001
        return _degraded("unsupported_document")

    # THE BOMB CHECK, AND IT IS BEFORE THE DECODE. `Image.open` is lazy: it reads the
    # header and nothing else, so `size` is known while the pixels are not yet
    # allocated. Checking here costs nothing; checking after `convert()` would be
    # checking after the allocation that was the whole danger.
    width, height = image.size
    if width <= 0 or height <= 0 or width * height > limits.max_image_pixels:
        return _degraded("unsupported_document", page_count=1)

    try:
        prepared = image.convert("L")
    except Exception:  # noqa: BLE001
        return _degraded("unsupported_document", page_count=1)

    per_page, confidence = _ocr_pages([prepared], limits)
    return _finish_ocr(
        [(1, texts) for texts in _normalize_pages(per_page)],
        page_count=1,
        confidence=confidence,
        limits=limits,
    )


def _extract_docx(data: bytes, limits: ExtractionLimits) -> ExtractionResult:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    try:
        document = Document(io.BytesIO(data))
    except Exception:  # noqa: BLE001 — PackageNotFoundError and friends
        return _degraded("unsupported_document")

    texts: list[str] = []

    def _cell_lines(table: Table) -> list[str]:
        out: list[str] = []
        for row in table.rows:
            for cell in row.cells:
                out.extend(_normalize_block(cell.text))
        return out

    # Contact details live in the HEADER as often as in the body on a templated résumé,
    # and a name or a phone number dropped here is a field RI-4 then has to ask for.
    for section in document.sections:
        for paragraph in section.header.paragraphs:
            texts.extend(_normalize_block(paragraph.text))

    # WALKING THE BODY XML, not `document.paragraphs` then `document.tables`. Those two
    # properties each return their own kind in order but give no interleaving, and
    # résumés are routinely laid out as a table — reading all paragraphs and then all
    # cells would shuffle a work history into an order nobody wrote.
    try:
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                texts.extend(_normalize_block(Paragraph(child, document).text))
            elif child.tag == qn("w:tbl"):
                texts.extend(_cell_lines(Table(child, document)))
    except Exception:  # noqa: BLE001
        return _degraded("unsupported_document")

    for section in document.sections:
        for paragraph in section.footer.paragraphs:
            texts.extend(_normalize_block(paragraph.text))

    if not texts:
        return _degraded("empty_document")

    # `page` stays None throughout: see `Line`. A .docx has no pages until it is
    # rendered, and this module renders nothing.
    lines, truncated = _number([(None, texts)], limits)
    if not lines:
        return _degraded("empty_document")
    return ExtractionResult(
        method="docx",
        lines=lines,
        page_count=None,
        ocr_confidence=None,
        degraded_reason=None,
        truncated=truncated,
    )


def _normalize_pages(per_page: list[list[str]]) -> list[list[str]]:
    """Tesseract's line strings through the same normalizer the other paths use."""
    return [[t for line in page for t in _normalize_block(line)] for page in per_page]


def _finish_ocr(
    blocks: list[tuple[int | None, list[str]]],
    *,
    page_count: int,
    confidence: float,
    limits: ExtractionLimits,
) -> ExtractionResult:
    lines, truncated = _number(blocks, limits)
    if not lines:
        return _degraded("empty_document", method="ocr", page_count=page_count)
    if confidence < limits.ocr_confidence_floor:
        # The score is KEPT on the degraded result on purpose. "How often, and by how
        # much, does OCR fall short?" is exactly the question RI-7 has to answer before
        # anyone claims this feature saves a worker time, and it is unanswerable if the
        # failures throw their measurement away.
        return _degraded(
            "ocr_below_floor", method="ocr", page_count=page_count, ocr_confidence=confidence
        )
    return ExtractionResult(
        method="ocr",
        lines=lines,
        page_count=page_count,
        ocr_confidence=confidence,
        degraded_reason=None,
        truncated=truncated,
    )


# ---------------------------------------------------------------------------
# The door
# ---------------------------------------------------------------------------


def extract(
    data: bytes, *, mime: str, limits: ExtractionLimits = DEFAULT_LIMITS
) -> ExtractionResult:
    """An uploaded résumé → citable lines, or a closed reason why not. Never raises."""
    declared = _MIME_TO_KIND.get(mime)
    if declared is None:
        return _degraded("unsupported_document")
    if not data:
        return _degraded("empty_document")

    sniffed = _sniff(data)
    if sniffed == "ole":
        # An OLE container where a .docx was promised is Office's encrypted wrapper.
        return _degraded("encrypted_document" if declared == "docx" else "unsupported_document")
    if sniffed != declared:
        # THE DECLARED TYPE IS A CLAIM. It reaches us from the client, through the
        # confirm call, and Supabase records whatever the upload said — so a JPEG
        # announced as a PDF would otherwise be handed to pypdf. Two agreeing sources
        # are not a guarantee when both derive from the same untrusted one.
        return _degraded("unsupported_document")

    if declared == "pdf":
        result = _extract_pdf(data, limits)
    elif declared == "docx":
        result = _extract_docx(data, limits)
    else:
        result = _extract_image(data, limits)

    # COUNTS AND A CONFIDENCE. No filename, no object key, no text, not one line of it.
    logger.info(
        "resume_import.extracted",
        extra={
            "extra": {
                "method": result.method,
                "line_count": len(result.lines),
                "page_count": result.page_count,
                "ocr_confidence": result.ocr_confidence,
                "degraded_reason": result.degraded_reason,
                "truncated": result.truncated,
            }
        },
    )
    return result
